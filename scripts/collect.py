"""Official-source collector. SQLite is authoritative; JSON is an atomic display cache."""
import argparse, calendar, csv, hashlib, io, json, os, pathlib, re, sqlite3, subprocess, time, urllib.request
from datetime import datetime, timezone
from urllib.parse import quote
ROOT=pathlib.Path(__file__).resolve().parents[1]
DATA=ROOT/'data'; DATA.mkdir(exist_ok=True)
REGISTRY=json.loads((ROOT/'lib/indicators.json').read_text(encoding='utf-8'))
WB_PAGE='https://www.worldbank.org/en/research/commodity-markets'
WB_FALLBACK='https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'
WB_COLUMNS=[
    ('NATGAS_US','Natural gas, US','mmbtu'),
    ('COAL_AUS','Coal, Australian','mt'),
    ('COPPER','Copper','mt'),
    ('ALUMINUM','Aluminum','mt'),
    ('IRON_ORE','Iron ore, cfr spot','dmtu'),
    ('NICKEL','Nickel','mt'),
    ('GOLD','Gold','troy oz'),
    ('SILVER','Silver','troy oz'),
]
WB_IDS=tuple(key for key,_,_ in WB_COLUMNS)
EIA_RETAIL_URL='https://www.eia.gov/electricity/data/state/xls/861m/HS861M%202010-.xlsx'
EIA_US_COMMERCIAL='EIA_US_COMMERCIAL'
# Daily market references supplement the World Bank monthly physical-price
# history. They remain separate series because contract basis and units differ.
MARKET_QUOTES={
    'MKT_COPPER':('HG=F','https://finance.yahoo.com/quote/HG%3DF/history/'),
    'MKT_ALUMINUM':('ALI=F','https://finance.yahoo.com/quote/ALI%3DF/history/'),
    'MKT_IRON_ORE':('TIO=F','https://finance.yahoo.com/quote/TIO%3DF/history/'),
    'MKT_GOLD':('GC=F','https://finance.yahoo.com/quote/GC%3DF/history/'),
    'MKT_SILVER':('SI=F','https://finance.yahoo.com/quote/SI%3DF/history/'),
}
MARKET_IDS=tuple(MARKET_QUOTES)
CREATE_NO_WINDOW=getattr(subprocess,'CREATE_NO_WINDOW',0) if os.name=='nt' else 0
WINDOWS_STARTUPINFO=None
if os.name=='nt':
    WINDOWS_STARTUPINFO=subprocess.STARTUPINFO();WINDOWS_STARTUPINFO.dwFlags|=subprocess.STARTF_USESHOWWINDOW;WINDOWS_STARTUPINFO.wShowWindow=subprocess.SW_HIDE
def now(): return datetime.now(timezone.utc).isoformat()
def download(url):
    result=subprocess.run(['curl.exe' if os.name=='nt' else 'curl','--fail','--location','--silent','--show-error','--max-time','65','--retry','2',url],capture_output=True,creationflags=CREATE_NO_WINDOW,startupinfo=WINDOWS_STARTUPINFO)
    if result.returncode: raise RuntimeError('Source fetch failed: '+result.stderr.decode(errors='replace')[-180:])
    return result.stdout
def download_market(url):
    """Use an in-process request for market references; Yahoo rejects curl clients."""
    request=urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0 (compatible; AI-Macro-Chain-Monitor/1.0)'})
    try:
        with urllib.request.urlopen(request,timeout=65) as response:return response.read()
    except Exception as exc:
        raise RuntimeError('Market reference fetch failed: '+str(exc)[-180:])
def parse_csv(raw):
    rows=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
    if not rows or rows[0][0] not in ('observation_date','DATE'): raise ValueError('Unexpected CSV header')
    result=[]; seen=set()
    for date,value in rows[1:]:
        datetime.strptime(date,'%Y-%m-%d')
        if date in seen: raise ValueError('Duplicate observation')
        seen.add(date)
        number=None if value in ('','.','NA') else float(value)
        if number is not None and not (-1e12<number<1e12): raise ValueError('Invalid numeric value')
        result.append({'date':date,'value':number})
    if not result: raise ValueError('Empty source')
    return sorted(result,key=lambda p:p['date'])
def parse_wb(raw):
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    rows=list(book['Monthly Prices'].values)
    header_index=next(i for i,row in enumerate(rows[:12]) if 'Copper' in row and 'Gold' in row and 'Silver' in row)
    columns={key:rows[header_index].index(name) for key,name,_ in WB_COLUMNS}
    units=rows[header_index+1]
    for key,_,expected in WB_COLUMNS:
        col=columns[key]
        if expected not in str(units[col]).lower(): raise ValueError('Unexpected commodity unit')
    result={key:[] for key in columns}
    for row in rows[header_index+2:]:
        if not re.fullmatch(r'\d{4}M\d{2}',str(row[0])): continue
        date=str(row[0]).replace('M','-')+'-01'
        for key,col in columns.items():
            val=row[col]; result[key].append({'date':date,'value':float(val) if isinstance(val,(int,float)) else None})
    if any(len(rows)<13 for rows in result.values()): raise ValueError('Commodity history too short')
    return result
def parse_eia_us_commercial(raw):
    from openpyxl import load_workbook
    rows=list(load_workbook(io.BytesIO(raw),read_only=True,data_only=True)['Monthly-States'].values)
    if rows[1][8]!='Revenue' or rows[1][9]!='Sales' or rows[2][8]!='Thousand Dollars' or rows[2][9]!='Megawatthours':raise ValueError('Unexpected EIA-861M commercial-electricity columns')
    grouped={}
    for row in rows[3:]:
        if not isinstance(row[0],(int,float)) or not isinstance(row[1],(int,float)) or not isinstance(row[2],str) or len(row[2])!=2:continue
        revenue,sales=row[8],row[9]
        if not isinstance(revenue,(int,float)) or not isinstance(sales,(int,float)) or sales<=0:continue
        grouped.setdefault((int(row[0]),int(row[1])),[]).append((row[2],float(revenue),float(sales)))
    points=[]
    for (year,month),states in sorted(grouped.items()):
        if len({state for state,_,_ in states})!=51:continue
        revenue=sum(value for _,value,_ in states);sales=sum(value for _,_,value in states)
        points.append({'date':f'{year}-{month:02}-{calendar.monthrange(year,month)[1]}','value':100*revenue/sales})
    if len(points)<13:raise ValueError('EIA-861M history too short or incomplete')
    return points
def parse_market_chart(raw):
    payload=json.loads(raw)
    chart=payload.get('chart',{});results=chart.get('result') or []
    if chart.get('error') or not results:raise ValueError('Yahoo Finance market chart unavailable')
    result=results[0];timestamps=result.get('timestamp') or [];closes=(result.get('indicators',{}).get('quote') or [{}])[0].get('close') or []
    if len(timestamps)!=len(closes) or not timestamps:raise ValueError('Yahoo Finance market chart has no usable close values')
    points=[];seen=set()
    for stamp,value in zip(timestamps,closes):
        date=datetime.fromtimestamp(stamp,timezone.utc).date().isoformat()
        if date in seen:continue
        seen.add(date);points.append({'date':date,'value':float(value) if isinstance(value,(int,float)) else None})
    if len(points)<13:raise ValueError('Yahoo Finance market history too short')
    return points
def connect():
    conn=sqlite3.connect(DATA/'observations.sqlite');conn.execute('PRAGMA journal_mode=WAL')
    conn.executescript('''CREATE TABLE IF NOT EXISTS observations(series_id TEXT, observation_date TEXT,value REAL,source_published_at TEXT,fetched_at TEXT,revision TEXT, PRIMARY KEY(series_id,observation_date,revision));
CREATE INDEX IF NOT EXISTS idx_observations_series_date ON observations(series_id,observation_date,fetched_at);
CREATE TABLE IF NOT EXISTS snapshots(series_id TEXT PRIMARY KEY,payload TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY,started_at TEXT,finished_at TEXT,status TEXT,detail TEXT);''')
    return conn
def save(conn,key,points,raw,source_url,started):
    digest=hashlib.sha256(raw).hexdigest(); version=DATA/'versions'/key;version.mkdir(parents=True,exist_ok=True)
    path=version/(digest+('.xlsx' if key in WB_IDS or key==EIA_US_COMMERCIAL else '.json' if key in MARKET_IDS else '.csv'))
    if not path.exists():path.write_bytes(raw)
    old=conn.execute('SELECT payload FROM snapshots WHERE series_id=?',(key,)).fetchone()
    old_payload=json.loads(old[0]) if old else {}; old_points={p['date']:p['value'] for p in old_payload.get('observations',[])}
    revision=started+'_'+digest[:10]
    changed=[(key,p['date'],p['value'],None,started,revision) for p in points if p['date'] not in old_points or old_points[p['date']]!=p['value']]
    conn.executemany('INSERT OR IGNORE INTO observations VALUES(?,?,?,?,?,?)',changed)
    payload={'id':key,'observations':points,'fetchedAt':started,'checkedAt':started,'sourcePublishedAt':None,'revision':revision if changed else old_payload.get('revision',revision),'sourceUrl':source_url,'status':'ready','error':None,'revisionCount':conn.execute('SELECT count(DISTINCT revision) FROM observations WHERE series_id=?',(key,)).fetchone()[0]}
    conn.execute('INSERT OR REPLACE INTO snapshots VALUES(?,?)',(key,json.dumps(payload,ensure_ascii=False)));conn.commit()
    return payload
def collect(import_dir=None, selected=None):
    started=now(); conn=connect(); failures=[]; wb=None; wb_raw=None; wb_url=WB_FALLBACK
    for config in REGISTRY:
        key=config['id'];ts=now()
        if selected is not None and key not in selected: continue
        try:
            if key in WB_IDS:
                if wb is None:
                    if import_dir: wb_raw=(pathlib.Path(import_dir)/'pink-sheet-monthly.xlsx').read_bytes()
                    else:
                        try:
                            html=download(WB_PAGE).decode(errors='replace');matches=re.findall(r'https://[^\s"<>]+CMO-Historical-Data-Monthly\.xlsx',html)
                            if matches:wb_url=matches[-1].replace('&amp;','&')
                        except Exception:pass
                        wb_raw=download(wb_url)
                    wb=parse_wb(wb_raw)
                raw=wb_raw;points=wb[key];url=wb_url
            elif key==EIA_US_COMMERCIAL:
                raw=download(EIA_RETAIL_URL);points=parse_eia_us_commercial(raw);url=EIA_RETAIL_URL
            elif key in MARKET_IDS:
                symbol,url=MARKET_QUOTES[key]
                raw=download_market('https://query2.finance.yahoo.com/v8/finance/chart/'+quote(symbol,safe='')+'?range=5y&interval=1d')
                points=parse_market_chart(raw)
            else:
                url='https://fred.stlouisfed.org/graph/fredgraph.csv?id='+key
                if import_dir:raw=(pathlib.Path(import_dir)/(key+'.csv')).read_bytes();points=parse_csv(raw)
                elif os.environ.get('FRED_API_KEY'):
                    api='https://api.stlouisfed.org/fred/series/observations?series_id='+key+'&file_type=json&api_key='+os.environ['FRED_API_KEY']
                    # Never log the URL containing the key.
                    try:
                        with urllib.request.urlopen(api,timeout=60) as response: payload=json.load(response)
                    except Exception:raise RuntimeError('FRED API fetch failed; credential withheld')
                    raw=('observation_date,'+key+'\n'+'\n'.join(p['date']+','+p['value'] for p in payload['observations'])).encode();points=parse_csv(raw)
                else:raw=download(url);points=parse_csv(raw)
            # A source must never move its latest observation backwards silently.
            previous=conn.execute('SELECT payload FROM snapshots WHERE series_id=?',(key,)).fetchone()
            if previous and json.loads(previous[0]).get('observations') and points[-1]['date']<json.loads(previous[0])['observations'][-1]['date']:raise ValueError('Source latest observation regressed')
            if import_dir:
                source_file=pathlib.Path(import_dir)/('pink-sheet-monthly.xlsx' if key in WB_IDS else 'eia-861m.xlsx' if key==EIA_US_COMMERCIAL else key+'.csv')
                ts=datetime.fromtimestamp(source_file.stat().st_mtime,timezone.utc).isoformat()
            save(conn,key,points,raw,url,ts);print(key,'OK',points[-1]['date'],flush=True)
        except Exception as exc:
            failures.append(key);old=conn.execute('SELECT payload FROM snapshots WHERE series_id=?',(key,)).fetchone()
            payload=json.loads(old[0]) if old else {'id':key,'observations':[],'fetchedAt':None,'sourcePublishedAt':None,'sourceUrl':config['release']}
            payload.update(status='cached' if payload['observations'] else 'fetch_failed',checkedAt=ts,error=str(exc))
            conn.execute('INSERT OR REPLACE INTO snapshots VALUES(?,?)',(key,json.dumps(payload,ensure_ascii=False)));conn.commit();print(key,'FAILED; retained cache',flush=True)
    series={key:json.loads(payload) for key,payload in conn.execute('SELECT series_id,payload FROM snapshots')}
    result={'generatedAt':now(),'series':series,'scheduler':{'mode':os.environ.get('MACRO_SCHEDULER','manual'),'lastRunAt':started,'failed':failures}}
    temp=DATA/'latest.tmp.json';temp.write_text(json.dumps(result,ensure_ascii=False,separators=(',',':')),encoding='utf-8');os.replace(temp,DATA/'latest.json')
    conn.execute('INSERT INTO runs(started_at,finished_at,status,detail) VALUES(?,?,?,?)',(started,now(),'partial_failure' if failures else 'success',json.dumps(failures)));conn.commit();conn.close()
    return result
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--import-dir');args=parser.parse_args();collect(args.import_dir)
