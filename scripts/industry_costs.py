"""World Bank monthly metals/gas and EIA state retail electricity: no price interpolation."""
import io,json,re,calendar
from openpyxl import load_workbook
from industry_common import DATA,now,fetch,definition,observation,persist
WB='https://www.worldbank.org/en/research/commodity-markets'
WB_FILE='https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'
EIA='https://www.eia.gov/electricity/data/state/xls/861m/HS861M%202010-.xlsx'
def parse_wb(raw):
 rows=list(load_workbook(io.BytesIO(raw),read_only=True,data_only=True)['Monthly Prices'].values);i=next(i for i,r in enumerate(rows[:12]) if 'Copper' in r);headers=rows[i];result={}
 for key,name in [('copper','Copper'),('aluminum','Aluminum'),('natural_gas','Natural gas, US')]:
  col=headers.index(name);pts=[]
  if ('mt' if key!='natural_gas' else 'mmbtu') not in str(rows[i+1][col]).lower():raise ValueError('Unexpected WB unit')
  for row in rows[i+2:]:
   if not re.fullmatch(r'\d{4}M\d{2}',str(row[0])):continue
   y,m=map(int,str(row[0]).split('M'));v=row[col];pts.append((f'{y}-{m:02}-{calendar.monthrange(y,m)[1]}',float(v) if isinstance(v,(int,float)) else None))
  result[key]=pts
 return result
def parse_eia(raw):
 book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);rows=list(book['Monthly-States'].values)
 assert rows[1][11]=='Price' and rows[2][11]=='Cents/kWh' and rows[1][15]=='Price'
 result={};national={}
 for row in rows[3:]:
  if isinstance(row[0],(int,float)) and isinstance(row[2],str) and len(row[2])==2:
   y,m=int(row[0]),int(row[1]);end=f'{y}-{m:02}-{calendar.monthrange(y,m)[1]}'
   for sector,revenue,sales in [('commercial',8,9),('industrial',12,13)]:
    national.setdefault((end,sector),[]).append((row[2],row[revenue],row[sales]))
  if not isinstance(row[0],(int,float)) or row[2] not in ('US','VA','TX','OH','GA','AZ'):continue
  y,m=int(row[0]),int(row[1]);end=f'{y}-{m:02}-{calendar.monthrange(y,m)[1]}'
  for sector,col in [('commercial',11),('industrial',15)]:
   key=str(row[2])+'.'+sector;v=row[col];result.setdefault(key,[]).append((end,float(v) if isinstance(v,(int,float)) and v>0 else None,str(row[3]),{'price_cents_kwh':v}))
 for (end,sector),points in national.items():
  value=100*sum(p[1] for p in points)/sum(p[2] for p in points) if len({p[0] for p in points})==51 and all(isinstance(p[1],(int,float)) and isinstance(p[2],(int,float)) for p in points) and sum(p[2] for p in points)>0 else None
  result.setdefault('US.'+sector,[]).append((end,value,'Calculated: 50 states + DC revenue-weighted by sales',{'revenue_thousand_usd':sum(p[1] for p in points) if value is not None else None,'sales_mwh':sum(p[2] for p in points) if value is not None else None,'region_count':len({p[0] for p in points})}))
 return {k:sorted(v) for k,v in result.items()}
def run():
 path=DATA/'costs.json';old=json.loads(path.read_text(encoding='utf-8')) if path.exists() else {'series':{}};result={'schemaVersion':'1','generatedAt':now(),'definitions':[],'series':{},'projects':[],'events':[]}
 try:
  url=WB_FILE
  try:
   html=fetch(WB)[0].decode(errors='replace');matches=re.findall(r'https://[^\s"<>]+CMO-Historical-Data-Monthly\.xlsx',html)
   if matches:url=matches[-1]
  except Exception:pass
  raw,stamp,version=fetch(url);metal=parse_wb(raw);error=None
 except Exception as e:metal={};error=str(e)
 for key,name,en in [('copper','铜','Copper'),('aluminum','铝','Aluminum'),('natural_gas','美国天然气','Natural gas, U.S. Henry Hub')]:
  id='WB.'+key;d=definition(id,name,en,'costs',key,'全球' if key!='natural_gas' else '美国','World Bank Pink Sheet',WB,'美元/公吨' if key!='natural_gas' else '美元/MMBtu','official','世界银行月度均价。天然气为美国 Henry Hub 基准；不是各地发电厂或数据中心采购合同。',frequency='monthly');result['definitions'].append(d)
  if key in metal:result['series'][id]={'observations':[observation(id,date,v,url,stamp,version) for date,v in metal[key]],'status':'ready','fetchedAt':stamp,'checkedAt':now()}
  else:result['series'][id]={**old['series'].get(id,{'observations':[]}),'status':'cached' if old['series'].get(id,{}).get('observations') else 'fetch_failed','checkedAt':now(),'error':error}
 try:raw,stamp,version=fetch(EIA);electricity=parse_eia(raw);error=None
 except Exception as e:electricity={};error=str(e)
 regions={'US':'美国全国','VA':'弗吉尼亚州','TX':'得克萨斯州','OH':'俄亥俄州','GA':'佐治亚州','AZ':'亚利桑那州'}
 for region,name in regions.items():
  for sector,label in [('commercial','商业'),('industrial','工业')]:
   key=region+'.'+sector;id='EIA.'+key;d=definition(id,name+label+'电价',name+' '+sector+' retail electricity price','power','electricity',region,'EIA-861M','https://www.eia.gov/electricity/data/state/','美分/kWh','official','终端零售收入/售电量的平均价格；商业/工业分别统计，非批发电价、非数据中心实际合同价格。月度初值可修订。',frequency='monthly');result['definitions'].append(d)
   if region=='US':d['valueType']='calculated';d['methodology']='美国50州及DC行业收入总和（千美元）/售电量总和（MWh）×100 = 美分/kWh；仅51个地区齐备时计算。非各州价格的简单平均，非数据中心合同价。'
   if key in electricity:result['series'][id]={'observations':[dict(observation(id,date,v,EIA,stamp,version,formula='sum(state revenue thousand USD) / sum(state sales MWh) * 100' if region=='US' else None,items=items),basis=status) for date,v,status,items in electricity[key]],'status':'ready','fetchedAt':stamp,'checkedAt':now()}
   else:result['series'][id]={**old['series'].get(id,{'observations':[]}),'status':'cached' if old['series'].get(id,{}).get('observations') else 'fetch_failed','checkedAt':now(),'error':error or '来源文件尚未取得该地区'}
 persist(result,path);print('industry costs:',sum(bool(s['observations']) for s in result['series'].values()),'connected')
if __name__=='__main__':run()
