"""ERCOT hourly system load aggregated to complete-month peak demand.

The annual archive is updated during the year.  We keep the hourly source at
its native Central Prevailing Time grain while parsing, and publish only a
deterministic monthly maximum for complete months.  Missing hours never become
zeros and an incomplete month never becomes an observation.
"""
from __future__ import annotations

import calendar
import concurrent.futures
import io
import json
import re
import struct
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urljoin
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from industry_common import DATA, definition, fetch, now, observation, persist


INDEX = "https://www.ercot.com/gridinfo/load/load_hist"
TIMEZONE = ZoneInfo("America/Chicago")
METRIC_ID = "POWER.load"
KNOWN_URLS = {
    2023: "https://www.ercot.com/files/docs/2023/02/09/Native_Load_2023.zip",
    2024: "https://www.ercot.com/files/docs/2024/02/06/Native_Load_2024.zip",
    2025: "https://www.ercot.com/files/docs/2025/02/11/Native_Load_2025.zip",
    2026: "https://www.ercot.com/files/docs/2026/02/10/Native_Load_2026.zip",
}
MONTHS = {name.lower(): number for number, name in enumerate(calendar.month_abbr) if name}
LOCAL_HEADER_FORMAT = "<4s5H3I2H"
LOCAL_HEADER_SIZE = 30


def _header(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _number(value: object) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _date_text(value: datetime) -> str:
    return value.strftime("%Y-%m-%d %H:%M")


def _parse_hour_ending(value: object) -> tuple[datetime, str, str]:
    """Return naive CPT wall time, the source label and its explicit DST tag."""
    raw = str(value).strip()
    dst_tag = "DST" if re.search(r"\bDST\b", raw, re.I) else ""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None), raw, dst_tag
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()), raw, dst_tag

    text = re.sub(r"\s+DST\b", "", raw, flags=re.I).strip()
    match_24 = re.fullmatch(r"(.+?)\s+24(?::00(?::00)?)?", text)
    if match_24:
        base, _, _ = _parse_hour_ending(match_24.group(1) + " 00:00")
        return base + timedelta(days=1), raw, dst_tag

    formats = (
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %H",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt), raw, dst_tag
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).replace(tzinfo=None), raw, dst_tag
    except ValueError as exc:
        raise ValueError(f"无法解析ERCOT Hour Ending: {raw[:80]}") from exc


def _read_zip_member(archive: bytes, filename: str) -> bytes:
    """Read a member even when ERCOT's central-directory sizes disagree."""
    zipped = ZipFile(io.BytesIO(archive))
    try:
        return zipped.read(filename)
    except BadZipFile:
        member = zipped.getinfo(filename)
        if zipped.fp is None:
            raise
        zipped.fp.seek(member.header_offset)
        header = zipped.fp.read(LOCAL_HEADER_SIZE)
        if len(header) != LOCAL_HEADER_SIZE:
            raise
        (
            signature,
            _version,
            flags,
            _compression,
            _time,
            _date,
            crc,
            compressed_size,
            uncompressed_size,
            _name_length,
            _extra_length,
        ) = struct.unpack(LOCAL_HEADER_FORMAT, header)
        if signature != b"PK\x03\x04" or flags & 0x08:
            raise
        member.CRC = crc
        member.compress_size = compressed_size
        member.file_size = uncompressed_size
        return zipped.read(filename)


def _workbook_bytes(raw: bytes) -> tuple[bytes, str]:
    archive = ZipFile(io.BytesIO(raw))
    members = [
        item.filename
        for item in archive.infolist()
        if not item.is_dir() and item.filename.lower().endswith((".xlsx", ".xlsm"))
    ]
    if members:
        preferred = next((name for name in members if "native_load" in name.lower()), members[0])
        return _read_zip_member(raw, preferred), preferred
    # A discovered link may point directly at the XLSX instead of an outer ZIP.
    if "[Content_Types].xml" in archive.namelist() and "xl/workbook.xml" in archive.namelist():
        return raw, "direct-xlsx"
    raise ValueError("ERCOT年度档案中没有可识别的XLSX工作簿")


def _expected_hours(year: int, month: int) -> int:
    start = datetime(year, month, 1, tzinfo=TIMEZONE)
    if month == 12:
        end = datetime(year + 1, 1, 1, tzinfo=TIMEZONE)
    else:
        end = datetime(year, month + 1, 1, tzinfo=TIMEZONE)
    return round((end.astimezone(timezone.utc) - start.astimezone(timezone.utc)).total_seconds() / 3600)


def parse_archive(raw: bytes, source_year: int) -> tuple[list[dict], dict]:
    workbook_raw, member_name = _workbook_bytes(raw)
    sheet = load_workbook(io.BytesIO(workbook_raw), read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header_row = None
    for candidate in rows:
        normalized = [_header(value) for value in candidate]
        if any(value in {"HOUR_ENDING", "HOURENDING", "HOUREND"} for value in normalized) and any(
            value in {"TOTAL", "ERCOT"} for value in normalized
        ):
            header_row = candidate
            break
    if header_row is None:
        raise ValueError("ERCOT工作簿未找到Hour Ending和TOTAL/ERCOT表头")

    headers = [_header(value) for value in header_row]
    time_index = next(i for i, value in enumerate(headers) if value in {"HOUR_ENDING", "HOURENDING", "HOUREND"})
    total_index = next(i for i, value in enumerate(headers) if value in {"TOTAL", "ERCOT"})
    zone_aliases = {
        "coast": {"COAST"},
        "east": {"EAST"},
        "far_west": {"FWEST", "FARWEST"},
        "north": {"NORTH"},
        "north_central": {"NCENT", "NORTHC"},
        "south": {"SOUTH", "SOUTHERN"},
        "south_central": {"SCENT", "SOUTHC"},
        "west": {"WEST"},
    }
    zone_indices: dict[str, int] = {}
    for name, aliases in zone_aliases.items():
        found = next((i for i, value in enumerate(headers) if value in aliases), None)
        if found is not None:
            zone_indices[name] = found
    if len(zone_indices) != 8:
        missing = sorted(set(zone_aliases) - set(zone_indices))
        raise ValueError("ERCOT天气分区列不完整: " + ", ".join(missing))

    by_month: dict[tuple[int, int], list[dict]] = defaultdict(list)
    invalid_by_month: Counter[tuple[int, int]] = Counter()
    for row in rows:
        if not row or time_index >= len(row) or row[time_index] in (None, ""):
            continue
        try:
            interval_end, raw_label, dst_tag = _parse_hour_ending(row[time_index])
        except ValueError:
            continue
        interval_start = interval_end - timedelta(hours=1)
        key = (interval_start.year, interval_start.month)
        if interval_start.year != source_year:
            continue
        total = _number(row[total_index] if total_index < len(row) else None)
        zones = [_number(row[index] if index < len(row) else None) for index in zone_indices.values()]
        if total is None or total <= 0 or any(value is None or value < 0 for value in zones):
            invalid_by_month[key] += 1
            continue
        zone_sum = sum(value for value in zones if value is not None)
        difference = abs(total - zone_sum)
        tolerance = max(1.0, abs(total) * 0.0001)
        by_month[key].append(
            {
                "total": total,
                "zone_difference": difference,
                "zone_sum_valid": difference <= tolerance,
                "interval_start": interval_start,
                "interval_end": interval_end,
                "raw_label": raw_label,
                "dst_tag": dst_tag,
            }
        )

    observations: list[dict] = []
    month_checks: dict[str, dict] = {}
    for (year, month), points in sorted(by_month.items()):
        expected = _expected_hours(year, month)
        actual = len(points)
        invalid = invalid_by_month[(year, month)]
        sum_failures = sum(not point["zone_sum_valid"] for point in points)
        wall_counts = Counter(point["interval_start"] for point in points)
        duplicate_wall_hours = sum(count - 1 for count in wall_counts.values())
        expected_duplicate_wall_hours = 1 if month == 11 else 0
        month_key = f"{year}-{month:02}"
        month_checks[month_key] = {
            "actualHours": actual,
            "expectedHours": expected,
            "invalidRows": invalid,
            "zoneSumFailures": sum_failures,
            "duplicateWallHours": duplicate_wall_hours,
            "expectedDuplicateWallHours": expected_duplicate_wall_hours,
        }
        if actual != expected or invalid or sum_failures or duplicate_wall_hours != expected_duplicate_wall_hours:
            continue
        peak = max(points, key=lambda point: point["total"])
        observations.append(
            {
                "periodStart": f"{year}-{month:02}-01",
                "periodEnd": f"{year}-{month:02}-{calendar.monthrange(year, month)[1]}",
                "value": peak["total"],
                "items": {
                    "source_year": source_year,
                    "source_member": member_name,
                    "complete_hour_count": actual,
                    "expected_hour_count": expected,
                    "peak_hour_ending_raw": peak["raw_label"],
                    "peak_interval_start_cpt": _date_text(peak["interval_start"]),
                    "peak_interval_end_cpt": _date_text(peak["interval_end"]),
                    "peak_dst_tag": peak["dst_tag"] or "none",
                    "peak_total_mw": peak["total"],
                    "peak_zone_sum_abs_diff_mw": peak["zone_difference"],
                    "month_max_zone_sum_abs_diff_mw": max(point["zone_difference"] for point in points),
                },
            }
        )
    return observations, {"sourceMember": member_name, "months": month_checks}


def discover_year_urls(raw: bytes, years: set[int]) -> dict[int, tuple[str, str | None]]:
    soup = BeautifulSoup(raw, "html.parser")
    result: dict[int, tuple[str, str | None]] = {}
    date_pattern = re.compile(r"\b([A-Z][a-z]{2})\s+(\d{1,2}),\s+(20\d\d)\b")
    for anchor in soup.find_all("a", href=True):
        href = anchor.get("href", "")
        text = " ".join(anchor.get_text(" ", strip=True).split())
        match_year = re.search(r"\b(20\d\d)\b", text)
        if not match_year or int(match_year.group(1)) not in years or not re.search(r"\.(?:zip|xlsx?)\b", href, re.I):
            continue
        year = int(match_year.group(1))
        context = text
        parent = anchor
        for _ in range(4):
            parent = parent.parent
            if parent is None:
                break
            context = " ".join(parent.get_text(" ", strip=True).split())
            if date_pattern.search(context):
                break
        published = None
        date_match = date_pattern.search(context)
        if date_match and date_match.group(1).lower() in MONTHS:
            month = MONTHS[date_match.group(1).lower()]
            published = f"{int(date_match.group(3)):04}-{month:02}-{int(date_match.group(2)):02}"
        result[year] = (urljoin(INDEX, href), published)
    return result


def _old_year_is_complete(old_observations: list[dict], year: int) -> bool:
    return len({point.get("periodEnd", "")[:7] for point in old_observations if point.get("periodEnd", "").startswith(f"{year}-")}) == 12


def _download_year(
    year: int, discovered: dict[int, tuple[str, str | None]]
) -> tuple[int, list[dict], dict, str, str, str, str | None]:
    candidates: list[tuple[str, str | None]] = []
    if year in discovered:
        candidates.append(discovered[year])
    if year in KNOWN_URLS and all(url != KNOWN_URLS[year] for url, _ in candidates):
        candidates.append((KNOWN_URLS[year], None))
    if not candidates:
        raise RuntimeError(f"{year}年ERCOT档案入口未发现，且没有已核验回退地址")
    failures = []
    for url, published in candidates:
        try:
            raw, stamp, version = fetch(url)
            parsed, checks = parse_archive(raw, year)
            if not parsed:
                raise ValueError(f"{year}年档案没有通过完整性校验的月份")
            return year, parsed, checks, url, stamp, version, published
        except Exception as exc:
            failures.append(str(exc))
    raise RuntimeError(f"{year}年ERCOT档案失败: " + " | ".join(failures))


def _next_check(latest_period: str | None) -> str:
    today_cpt = datetime.now(TIMEZONE).date()
    first_this_month = today_cpt.replace(day=1)
    expected_end = first_this_month - timedelta(days=1)
    if latest_period and latest_period >= expected_end.isoformat():
        if first_this_month.month == 12:
            next_month = date(first_this_month.year + 1, 1, 10)
        else:
            next_month = date(first_this_month.year, first_this_month.month + 1, 10)
        check = datetime.combine(next_month, datetime.min.time(), TIMEZONE) + timedelta(hours=12)
    else:
        check = datetime.now(TIMEZONE) + timedelta(days=7)
    return check.astimezone(timezone.utc).isoformat()


def run() -> None:
    path = DATA / "power-load.json"
    old = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"series": {}}
    old_series = old.get("series", {}).get(METRIC_ID, {})
    old_observations = old_series.get("observations", [])
    current_year = datetime.now(TIMEZONE).year
    years = list(range(current_year - 3, current_year + 1))

    next_check_raw = old_series.get("nextCheckAt")
    defer_download = False
    if old_observations and next_check_raw:
        try:
            defer_download = datetime.fromisoformat(next_check_raw) > datetime.now(timezone.utc)
        except ValueError:
            pass

    discovered: dict[int, tuple[str, str | None]] = {}
    index_error = old_series.get("sourceIndexError") if defer_download else None
    if not defer_download:
        try:
            index_raw, _, _ = fetch(INDEX)
            discovered = discover_year_urls(index_raw, set(years))
        except Exception as exc:
            index_error = str(exc)

    # Completed historical years are immutable enough for this product: reuse
    # their persisted monthly facts and avoid repeatedly downloading a report.
    years_to_fetch = [] if defer_download else [
        year for year in years if year == current_year or not _old_year_is_complete(old_observations, year)
    ]
    results = []
    failures: list[str] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=min(4, max(1, len(years_to_fetch)))) as pool:
        futures = {pool.submit(_download_year, year, discovered): year for year in years_to_fetch}
        for future in concurrent.futures.as_completed(futures):
            year = futures[future]
            try:
                results.append(future.result())
            except Exception as exc:
                failures.append(f"{year}: {exc}")

    merged = {point["periodEnd"]: point for point in old_observations}
    source_checks: dict[str, dict] = {
        year: checks
        for year, checks in old_series.get("sourceChecks", {}).items()
        if year.isdigit() and int(year) in years
    }
    for year, parsed, checks, url, stamp, version, published in results:
        source_checks[str(year)] = {"url": url, **checks}
        for item in parsed:
            point = observation(
                METRIC_ID,
                item["periodEnd"],
                item["value"],
                url,
                stamp,
                version,
                start=item["periodStart"],
                formula="max(ERCOT hourly TOTAL MW) for all complete CPT intervals in the calendar month",
                items=item["items"],
                published=published,
            )
            point["basis"] = "ERCOT actual hourly system load; complete-month maximum"
            merged[point["periodEnd"]] = point

    observations = sorted(merged.values(), key=lambda point: point["periodEnd"])
    if len(observations) > 48:
        observations = observations[-48:]
    if len(observations) < 36:
        failures.append(f"完整月仅{len(observations)}个，少于所需36个")
    if not results and index_error:
        failures.append("年度入口页: " + index_error)
    error = old_series.get("error") if defer_download else "; ".join(dict.fromkeys(failures)) or None
    status = old_series.get("status", "cached") if defer_download else "ready" if not error else "cached" if observations else "fetch_failed"

    metric = definition(
        METRIC_ID,
        "ERCOT月度实际峰值负荷",
        "ERCOT monthly actual peak system load",
        "power",
        "load",
        "ERCOT",
        "Electric Reliability Council of Texas (ERCOT)",
        INDEX,
        "MW",
        "calculated",
        "对ERCOT公开小时实际负荷TOTAL/ERCOT列，在Central Prevailing Time自然月内仅对完整月份取最大值。Hour Ending先换算为interval end，再减一小时得到interval start并据此归月；月末HE24不会错放到下月。缺失小时、分区合计异常或部分月份不发布，不填零、不插值。",
        eligible=False,
        frequency="monthly",
        export=True,
    )
    metric.update(
        {
            "sourceAdapter": "power-load",
            "aggregation": "none",
            "normalUpdateDelayDays": 45,
            "isComparableAcrossEntities": False,
            "interpretation": "月度峰值反映ERCOT系统在最紧张小时的实际用电水平，天气、人口、工业活动和其他负荷都会影响它。",
            "transmission": "峰值持续抬升可能增加电源、输配电和备用需求，但总负荷不能单独归因于数据中心或AI建设。",
            "crossCheck": "同时观察天气、ERCOT长期负荷预测、区域电价、数据中心并网和项目投运证据。",
            "licenseNote": "仅保存并导出派生月度峰值及来源链接，不镜像原始工作簿；使用时保留ERCOT来源、版权与Terms of Use提示。",
        }
    )
    latest = observations[-1]["periodEnd"] if observations else None
    result = {
        "schemaVersion": "1",
        "generatedAt": now(),
        "definitions": [metric],
        "series": {
            METRIC_ID: {
                "observations": observations,
                "status": status,
                "fetchedAt": observations[-1]["fetchedAt"] if observations else old_series.get("fetchedAt"),
                "checkedAt": now(),
                "nextCheckAt": _next_check(latest),
                "error": error,
                "note": f"最新通过完整性校验的月份：{latest or '无'}；共{len(observations)}个月。ERCOT总负荷不等于数据中心负荷。",
                "sourceChecks": source_checks,
                "sourceIndexError": index_error,
            }
        },
        "projects": [],
        "events": [],
    }
    persist(result, path)
    print(f"ERCOT monthly peaks: {len(observations)} complete months; status={status}; latest={latest or 'none'}")
    if error:
        print("ERCOT load warning:", error)


if __name__ == "__main__":
    run()
